import openpyxl
from decimal import Decimal
from typing import Dict, Any, Optional
import logging
from pathlib import Path

from src.db.connection import DatabaseConnection
from src.db.queries import QueryRepository
from src.answer.models import (
    MasterAnswerData,
    InventoryAnswer,
    FixedAssetAnswer,
    FixedAssetItem,
    GeneralBalanceAnswer,
    AccountBalanceItem,
    TransactionAnswer,
    TransactionItem,
    FinancialReportAnswer,
    AccountingGraph,
)
from src.grading.accounting_graph import AccountingGraphBuilder

logger = logging.getLogger(__name__)


class AnswerLoader:
    """Load đáp án từ SQL Server hoặc Excel fallback."""

    def __init__(
        self,
        db_conn: DatabaseConnection,
        query_repo: QueryRepository,
        config: Dict[str, Any],
    ):
        self.db_conn = db_conn
        self.query_repo = query_repo
        self.config = config
        self.answer_config = config.get("answer", {})
        self.answer_db = self.answer_config.get("database", "HungBinh2024")
        self.excel_path = self.answer_config.get("excel_fallback", "d12.xlsm")
        self.excel_sheet = self.answer_config.get(
            "excel_sheet_answer", "DapAn"
        )
        self.graph_builder = AccountingGraphBuilder()

    def load(self) -> MasterAnswerData:
        """Load đáp án tổng hợp và đồ thị nghiệp vụ cấp dòng."""
        logger.info("Đang load đáp án từ database [%s]...", self.answer_db)
        master = MasterAnswerData()

        try:
            db_inv = self.query_repo.execute(
                self.db_conn, self.answer_db, "inventory_balance"
            )
            if db_inv and db_inv[0]["item_count"] > 0:
                master.inventory = InventoryAnswer(
                    item_count=int(db_inv[0]["item_count"]),
                    total_qty=Decimal(str(db_inv[0]["total_qty"])),
                    total_amount=Decimal(str(db_inv[0]["total_amount"])),
                )

            db_fa = self.query_repo.execute(
                self.db_conn, self.answer_db, "fixed_asset_balance"
            )
            if db_fa:
                fa_items = []
                for row in db_fa:
                    fa_items.append(
                        FixedAssetItem(
                            code=str(row.get("FixedAssetCode", "")),
                            name=str(row.get("FixedAssetName", "")),
                            org_price=Decimal(str(row.get("OrgPrice", 0))),
                            depreciation_amount=Decimal(
                                str(row.get("DepreciationAmount", 0))
                            ),
                            accum_depreciation_amount=Decimal(
                                str(row.get("AccumDepreciationAmount", 0))
                            ),
                            lifetime_months=int(
                                row.get("LifeTimeInMonth", 0) or 0
                            ),
                            remaining_months=int(
                                row.get("LifeTimeRemainingInMonth", 0) or 0
                            ),
                        )
                    )
                master.fixed_asset = FixedAssetAnswer(items=fa_items)

            db_gb = self.query_repo.execute(
                self.db_conn, self.answer_db, "general_balance"
            )
            if db_gb:
                gb_accounts = {}
                for row in db_gb:
                    code = str(row["AccountCode"])
                    gb_accounts[code] = AccountBalanceItem(
                        account_code=code,
                        debit_amount=Decimal(
                            str(row.get("DebitAmount", 0))
                        ),
                        debit_amount_oc=Decimal(
                            str(row.get("DebitAmountOC", 0))
                        ),
                        credit_amount=Decimal(
                            str(row.get("CreditAmount", 0))
                        ),
                        credit_amount_oc=Decimal(
                            str(row.get("CreditAmountOC", 0))
                        ),
                        quantity=Decimal(str(row.get("Quantity", 0))),
                    )
                master.general_balance = GeneralBalanceAnswer(
                    accounts=gb_accounts
                )

            db_tx = self.query_repo.execute(
                self.db_conn, self.answer_db, "transactions"
            )
            if db_tx:
                tx_entries = {}
                for row in db_tx:
                    task_id = int(row["TaskID"])
                    code = str(row["AccountCode"])
                    tx_entries[(task_id, code)] = TransactionItem(
                        task_id=task_id,
                        account_code=code,
                        amount=Decimal(str(row.get("Amount", 0))),
                        amount_oc=Decimal(str(row.get("AmountOC", 0))),
                        quantity=Decimal(str(row.get("Quantity", 0))),
                    )
                master.transactions = TransactionAnswer(entries=tx_entries)

            db_fr = self.query_repo.execute(
                self.db_conn, self.answer_db, "financial_reports"
            )
            if db_fr:
                fr_items = {}
                for row in db_fr:
                    report_type = str(row["ReportType"])
                    code = str(row["ItemCode"])
                    fr_items[(report_type, code)] = Decimal(
                        str(row.get("Amount", 0))
                    )
                master.financial_reports = FinancialReportAnswer(
                    items=fr_items
                )
        except Exception as exc:
            logger.warning(
                "Không thể load đầy đủ đáp án tổng hợp từ DB [%s]: %s",
                self.answer_db,
                exc,
            )

        # Đồ thị cấp dòng được load độc lập để một query tổng hợp lỗi không làm
        # mất toàn bộ dữ liệu workflow.
        master.accounting_graph = self._load_accounting_graph()

        if Path(self.excel_path).exists():
            logger.info(
                "Bổ sung đáp án từ file Excel: %s (sheet '%s')",
                self.excel_path,
                self.excel_sheet,
            )
            self._fill_from_excel(master)

        try:
            from src.db.sqlite_storage import SQLiteStorage

            storage = SQLiteStorage()
            if master.transactions or master.inventory:
                storage.save_master_answer(master, self.answer_db)
        except Exception as exc:
            logger.debug(
                "Không thể lưu master answer cache vào SQLite: %s", exc
            )

        return master

    def _load_accounting_graph(self) -> Optional[AccountingGraph]:
        raw: Dict[str, Any] = {}
        successful_queries = 0
        for query_id in AccountingGraphBuilder.QUERY_IDS:
            try:
                raw[query_id] = self.query_repo.execute(
                    self.db_conn, self.answer_db, query_id
                )
                successful_queries += 1
            except Exception as exc:
                logger.warning(
                    "Không load được query workflow [%s] từ đáp án [%s]: %s",
                    query_id,
                    self.answer_db,
                    exc,
                )
                raw[query_id] = []

        if successful_queries == 0:
            return None
        graph = self.graph_builder.build(raw)
        if not (
            graph.entities
            or graph.opening_balances
            or graph.events
            or graph.inventory_ledger
            or graph.general_ledger
        ):
            return None
        return graph

    def _fill_from_excel(self, master: MasterAnswerData):
        """Đọc sheet DapAn nếu DB thiếu dữ liệu tổng hợp cũ."""
        wb = openpyxl.load_workbook(
            self.excel_path, read_only=True, data_only=True
        )
        if self.excel_sheet not in wb.sheetnames:
            wb.close()
            return

        ws = wb[self.excel_sheet]

        if not master.inventory or master.inventory.item_count == 0:
            a8 = ws["A8"].value
            b8 = ws["B8"].value
            c8 = ws["C8"].value
            if a8 is not None and b8 is not None and c8 is not None:
                master.inventory = InventoryAnswer(
                    item_count=int(a8),
                    total_qty=Decimal(str(b8)),
                    total_amount=Decimal(str(c8)),
                )

        if not master.fixed_asset or not master.fixed_asset.items:
            fa_items = []
            for row_number in range(8, 15):
                original_price = ws[f"D{row_number}"].value
                if original_price is not None and isinstance(
                    original_price, (int, float)
                ):
                    fa_items.append(
                        FixedAssetItem(
                            code=f"TS00{row_number - 7}",
                            name=f"Tài sản cố định {row_number - 7}",
                            org_price=Decimal(str(original_price)),
                            depreciation_amount=Decimal(
                                str(ws[f"E{row_number}"].value or 0)
                            ),
                            accum_depreciation_amount=Decimal(
                                str(ws[f"F{row_number}"].value or 0)
                            ),
                            lifetime_months=int(
                                ws[f"G{row_number}"].value or 0
                            ),
                            remaining_months=int(
                                ws[f"H{row_number}"].value or 0
                            ),
                        )
                    )
            if fa_items:
                master.fixed_asset = FixedAssetAnswer(items=fa_items)

        if not master.general_balance or not master.general_balance.accounts:
            gb_accounts = {}
            for row_number in range(8, 30):
                account = ws[f"I{row_number}"].value
                if account:
                    code = str(account).strip()
                    gb_accounts[code] = AccountBalanceItem(
                        account_code=code,
                        debit_amount=Decimal(
                            str(ws[f"J{row_number}"].value or 0)
                        ),
                        debit_amount_oc=Decimal(
                            str(ws[f"K{row_number}"].value or 0)
                        ),
                        credit_amount=Decimal(
                            str(ws[f"L{row_number}"].value or 0)
                        ),
                        credit_amount_oc=Decimal(
                            str(ws[f"M{row_number}"].value or 0)
                        ),
                        quantity=Decimal(
                            str(ws[f"N{row_number}"].value or 0)
                        ),
                    )
            if gb_accounts:
                master.general_balance = GeneralBalanceAnswer(
                    accounts=gb_accounts
                )

        if not master.transactions or not master.transactions.entries:
            tx_entries = {}
            for row_number in range(8, 45):
                task = ws[f"O{row_number}"].value
                account = ws[f"P{row_number}"].value
                if task is not None and account is not None:
                    try:
                        task_id = int(task)
                        code = str(account).strip()
                        tx_entries[(task_id, code)] = TransactionItem(
                            task_id=task_id,
                            account_code=code,
                            amount=Decimal(
                                str(ws[f"Q{row_number}"].value or 0)
                            ),
                            amount_oc=Decimal(
                                str(ws[f"R{row_number}"].value or 0)
                            ),
                            quantity=Decimal(
                                str(ws[f"S{row_number}"].value or 0)
                            ),
                        )
                    except (ValueError, TypeError):
                        pass
            if tx_entries:
                master.transactions = TransactionAnswer(entries=tx_entries)

        if (
            not master.financial_reports
            or not master.financial_reports.items
        ):
            fr_items = {}
            for row_number in range(8, 40):
                report_type = ws[f"T{row_number}"].value
                item_code = ws[f"U{row_number}"].value
                amount = ws[f"V{row_number}"].value
                if (
                    report_type is not None
                    and item_code is not None
                    and amount is not None
                ):
                    fr_items[
                        (str(report_type).strip(), str(item_code).strip())
                    ] = Decimal(str(amount))
            if fr_items:
                master.financial_reports = FinancialReportAnswer(
                    items=fr_items
                )

        wb.close()
