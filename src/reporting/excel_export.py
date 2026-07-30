import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any
from pathlib import Path
from datetime import datetime
import logging
from src.grading.engine import StudentGradeReport

logger = logging.getLogger(__name__)

class ExcelExporter:
    """Xuất bảng điểm tổng hợp của tất cả sinh viên ra file Excel."""

    def __init__(self, output_dir: str = "output"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export(self, reports: List[StudentGradeReport], filename: str = None) -> str:
        if not filename:
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"BangDiem_HTTTKT_{ts}.xlsx"

        file_path = self.output_dir / filename

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "BangDiem"

        # Headers & Styling
        title_font = Font(name="Arial", size=14, bold=True, color="1F4E78")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")

        # Title Block
        ws.merge_cells("A1:K1")
        ws["A1"] = "BẢNG ĐIỂM CHẤM THI HỆ THỐNG THÔNG TIN KẾ TOÁN (MISA SME.NET)"
        ws["A1"].font = title_font
        ws["A1"].alignment = center_align

        ws["A2"] = f"Thời gian xuất báo cáo: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
        ws["A2"].font = Font(name="Arial", size=10, italic=True)

        # Table Header Row
        headers = [
            "STT", "Mã SV", "Họ", "Tên", "Số máy",
            "Số dư HTK (10)", "Số dư TSCĐ (5)", "Số dư Sổ cái (25)", "Nghiệp vụ (55)", "BCTC (5)",
            "Tổng điểm (10)"
        ]

        row_num = 5
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=row_num, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        # Data Rows
        for idx, r in enumerate(reports, 1):
            row_num += 1
            st = r.student
            scores = r.score_data.get("module_results", {})

            inv_s = scores.get("inventory_balance", None)
            fa_s = scores.get("fixed_asset_balance", None)
            gb_s = scores.get("general_balance", None)
            tx_s = scores.get("transactions", None)
            fr_s = scores.get("financial_reports", None)

            inv_val = inv_s.score if inv_s else 0.0
            fa_val = fa_s.score if fa_s else 0.0
            gb_val = gb_s.score if gb_s else 0.0
            tx_val = tx_s.score if tx_s else 0.0
            fr_val = fr_s.score if fr_s else 0.0
            final_val = r.score_data.get("final_score", 0.0)

            vals = [
                idx, st.student_id, st.last_name, st.first_name, st.computer_name or "N/A",
                inv_val, fa_val, gb_val, tx_val, fr_val, final_val
            ]

            for col_num, val in enumerate(vals, 1):
                cell = ws.cell(row=row_num, column=col_num, value=val)
                cell.border = thin_border
                cell.font = Font(name="Arial", size=10)
                if col_num in [1, 2, 5]:
                    cell.alignment = center_align
                elif col_num in [3, 4]:
                    cell.alignment = left_align
                else:
                    cell.alignment = right_align
                    if isinstance(val, (int, float)):
                        cell.number_format = "0.00"

                # Highlight final score
                if col_num == 11:
                    cell.font = Font(name="Arial", size=10, bold=True, color="C00000")

        # Auto-fit column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(file_path)
        logger.info(f"Đã xuất bảng điểm thành công ra file: {file_path}")
        return str(file_path)
